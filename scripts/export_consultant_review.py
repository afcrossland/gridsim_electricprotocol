#!/usr/bin/env python3
"""Build a consultant-review spreadsheet: one tab per jurisdiction, one row
per question, with a dropdown of that question's own rubric answers and the
evidence already on file.

Two-step build, same reasoning as the other scripts in this repo (see
README's "House style" / the reuse note in extractCountryData.ts): the
answer-merging logic (which sourced answer wins for a given country and
question) lives in TypeScript and should stay there rather than being
re-derived here, so this script never computes an answer itself - it only
reads what the real app already computed.

    1. npx vitest run --config scripts/vitest.extract.config.ts
       -> scripts/extractCountryData.ts imports the app's own
          sourcedResponses()/sourcedCountries() and writes
          scripts/_countryData.json (gitignored, regenerate each run - it's
          a snapshot of live-computed data, not a source of truth itself).
    2. python3 scripts/export_consultant_review.py
       -> reads that JSON, writes Solar Policy Explorer - Consultant
          Review.xlsx to the repo root.

Layout per tab: column A question area (section), B question, C answer
(data-validation dropdown built from that question's own rubric, pre-filled
with the current answer where one exists), D-F evidence title/source/notes.
A question with more than one evidence entry stacks them in the same three
cells, separated by a blank line, rather than repeating the row - see the
README/skill note on this being a deliberate choice, not the only option.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA_PATH = Path(__file__).resolve().parent / "_countryData.json"
OUT_PATH = ROOT / "Solar Policy Explorer - Consultant Review.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="00ABBB")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
HEADERS = ["Question area", "Question", "Answer", "Evidence title", "Evidence source", "Evidence notes"]
COL_WIDTHS = [26, 46, 34, 26, 34, 46]


def safe_sheet_name(name: str, used: set[str]) -> str:
    """Excel sheet names: <=31 chars, no : \\ / ? * [ ], and must be unique."""
    cleaned = "".join(c for c in name if c not in ':\\/?*[]')[:31]
    candidate = cleaned or "Jurisdiction"
    n = 2
    while candidate in used:
        suffix = f" {n}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def build_rubric_lists_sheet(wb: Workbook, questions: list[dict]) -> dict[str, str]:
    """One column per question on a hidden sheet, holding its rubric labels -
    a dropdown's source range has to be a real range of cells, not an
    inline list, once the combined label text is long (Excel's inline-list
    data validation caps out around 255 characters, and several of this
    protocol's rubric tiers are full sentences). Returns questionId -> the
    A1 range string covering its labels.
    """
    sheet = wb.create_sheet("RubricLists")
    sheet.sheet_state = "hidden"
    ranges: dict[str, str] = {}
    for i, q in enumerate(questions, start=1):
        col = get_column_letter(i)
        sheet.cell(row=1, column=i, value=q["id"])
        for r, tier in enumerate(q["rubric"], start=2):
            sheet.cell(row=r, column=i, value=tier["label"])
        last_row = 1 + len(q["rubric"])
        ranges[q["id"]] = f"RubricLists!${col}$2:${col}${last_row}"
    return ranges


def main() -> None:
    payload = json.loads(DATA_PATH.read_text())
    sections = {s["id"]: s["title"] for s in payload["sections"]}
    questions = payload["questions"]
    countries = payload["countries"]

    wb = Workbook()
    wb.remove(wb.active)  # default blank sheet - every tab below is a real one

    rubric_ranges = build_rubric_lists_sheet(wb, questions)

    used_names: set[str] = set()
    # Most-complete first - a consultant reviewing this gets the most
    # substantive tabs up front rather than paging past dozens of
    # near-empty ones to find something worth checking.
    for code in sorted(countries, key=lambda c: countries[c]["completeness"], reverse=True):
        country = countries[code]
        sheet_name = safe_sheet_name(f"{country['name']} ({code})", used_names)
        ws = wb.create_sheet(sheet_name)

        for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

        row = 2
        for q in questions:
            answer = country["answers"].get(q["id"])
            rubric_by_score = {t["score"]: t["label"] for t in q["rubric"]}

            ws.cell(row=row, column=1, value=sections.get(q["sectionId"], "")).alignment = WRAP
            ws.cell(row=row, column=2, value=q["text"]).alignment = WRAP

            answer_cell = ws.cell(row=row, column=3)
            answer_cell.alignment = WRAP
            if answer is not None:
                answer_cell.value = rubric_by_score.get(answer["score"], "")

            dv = DataValidation(type="list", formula1=rubric_ranges[q["id"]], allow_blank=True)
            dv.error = "Pick one of this question's own answer options from the dropdown."
            dv.errorTitle = "Not a valid answer for this question"
            ws.add_data_validation(dv)
            dv.add(answer_cell)

            evidence = answer["evidence"] if answer else []
            title_cell = ws.cell(row=row, column=4, value="\n\n".join(e["title"] for e in evidence if e["title"]))
            source_cell = ws.cell(row=row, column=5, value="\n\n".join(e["source"] for e in evidence if e["source"]))
            notes_cell = ws.cell(row=row, column=6, value="\n\n".join(e["note"] for e in evidence if e["note"]))
            for cell in (title_cell, source_cell, notes_cell):
                cell.alignment = WRAP

            row += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} - {len(countries)} tabs, {len(questions)} rows each")


if __name__ == "__main__":
    main()
