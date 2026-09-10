#!/usr/bin/env python3
"""Build a Word doc listing every Electric Protocol question and its
possible answers, styled with the same GSC colours the app itself uses
(src/mui-theme.tsx) - reference material for anyone reading the
questionnaire outside the app (a consultant, a new researcher).

Reads src/data/protocol.seed.json directly rather than going through the
app's TS logic - unlike the consultant-review spreadsheet, this doesn't need
any per-country answer/merge data, just the static question set, so there's
nothing here that could drift from what sourcedResponses() would compute and
a plain read is the simpler, equally-correct choice (see the note on when to
reach for the Vitest-runner trick in scripts/extractCountryData.ts instead).

Usage: python3 scripts/export_questions_doc.py
"""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
SEED_PATH = ROOT / "src" / "data" / "protocol.seed.json"
OUT_PATH = ROOT / "Electric Protocol - Questions and Answers.docx"

# GSC brand colours - same values as src/mui-theme.tsx.
AQUA = RGBColor(0x00, 0xAB, 0xBB)
TEAL = RGBColor(0x00, 0x81, 0x94)
CITRUS = RGBColor(0xFB, 0xB1, 0x14)
DEEP_GRAY = RGBColor(0x3B, 0x38, 0x38)
TEXT_SECONDARY = RGBColor(0x6B, 0x72, 0x80)
OFF_WHITE_HEX = "F4F1E9"
AQUA_HEX = "00ABBB"

# impactColor()'s gradient, src/lib/scoring.ts - identical constants, so a
# question's weight badge here matches its "Impactfullness" chip in the app.
MAX_IMPACT = 5
IMPACT_LOW_RGB = (255, 243, 74)  # GSC Bright Yellow
IMPACT_HIGH_RGB = (46, 125, 50)


def impact_color(weight: float) -> RGBColor:
    t = min(1.0, max(0.0, weight / MAX_IMPACT))
    r, g, b = (round(c + (IMPACT_HIGH_RGB[i] - c) * t) for i, c in enumerate(IMPACT_LOW_RGB))
    return RGBColor(r, g, b)


def impact_text_color(weight: float) -> RGBColor:
    t = min(1.0, max(0.0, weight / MAX_IMPACT))
    return RGBColor(0xFF, 0xFF, 0xFF) if t > 0.5 else DEEP_GRAY


def set_cell_background(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_run(paragraph, text: str, *, color=None, bold=False, italic=False, size=None, font="Calibri"):
    run = paragraph.add_run(text)
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    return run


def build() -> None:
    seed = json.loads(SEED_PATH.read_text())
    sections = sorted(seed["sections"], key=lambda s: s["order"])
    questions_by_section: dict[str, list[dict]] = {}
    for q in sorted(seed["questions"], key=lambda q: q["order"]):
        questions_by_section.setdefault(q["sectionId"], []).append(q)

    doc = Document()
    for style_name in ("Normal",):
        doc.styles[style_name].font.name = "Calibri"
        doc.styles[style_name].font.size = Pt(10.5)
        doc.styles[style_name].font.color.rgb = DEEP_GRAY

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, seed.get("title", "Electric Protocol"), color=TEAL, bold=True, size=28)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(subtitle, "Questions and possible answers", color=CITRUS, italic=True, bold=True, size=14)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(
        note,
        f"{len(seed['questions'])} questions across {len(sections)} sections - Solar Policy Explorer reference",
        color=TEXT_SECONDARY,
        size=10,
    )
    doc.add_page_break()

    for section in sections:
        # Aqua bar behind the section title - same colour as the app's own
        # primary.main, done via a one-cell borderless table rather than
        # python-docx's more awkward raw-paragraph-shading API.
        bar = doc.add_table(rows=1, cols=1)
        bar.alignment = WD_TABLE_ALIGNMENT.LEFT
        bar.autofit = True
        cell = bar.cell(0, 0)
        cell.width = Inches(6.5)
        p = cell.paragraphs[0]
        add_run(p, section["title"], color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, size=13)
        set_cell_background(cell, AQUA_HEX)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

        for q in questions_by_section.get(section["id"], []):
            q_para = doc.add_paragraph()
            q_para.paragraph_format.space_before = Pt(10)
            q_para.paragraph_format.space_after = Pt(2)
            if q.get("subsection"):
                add_run(q_para, f"{q['subsection']} - ", color=TEXT_SECONDARY, italic=True, size=9.5)
            add_run(q_para, q["text"], color=DEEP_GRAY, bold=True, size=11)

            weight = q["weight"]
            impact_para = doc.add_paragraph()
            impact_para.paragraph_format.space_after = Pt(4)
            badge_run = add_run(
                impact_para,
                f"  Impactfullness: {weight:g}  ",
                color=impact_text_color(weight),
                bold=True,
                size=8.5,
            )
            # A 1x1 table per question just for a coloured badge would be
            # overkill - python-docx has no highlight-colour helper beyond a
            # fixed palette, so shade the run's own background directly via
            # its rPr instead, which Word does support.
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "%02X%02X%02X" % (impact_color(weight)[0], impact_color(weight)[1], impact_color(weight)[2]))
            badge_run._r.get_or_add_rPr().append(shd)

            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.autofit = False
            hdr = table.rows[0].cells
            hdr[0].width = Inches(0.7)
            hdr[1].width = Inches(5.6)
            for cell, label in zip(hdr, ("Score", "Possible answer")):
                cell.text = ""
                p = cell.paragraphs[0]
                add_run(p, label, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, size=9.5)
                set_cell_background(cell, TEAL.__str__())

            for i, tier in enumerate(sorted(q["rubric"], key=lambda t: t["score"])):
                row = table.add_row().cells
                row[0].width = Inches(0.7)
                row[1].width = Inches(5.6)
                score_p = row[0].paragraphs[0]
                add_run(score_p, str(tier["score"]), color=DEEP_GRAY, bold=True, size=10)
                score_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                answer_p = row[1].paragraphs[0]
                add_run(answer_p, tier["label"], color=DEEP_GRAY, size=10)
                if i % 2 == 1:
                    for cell in row:
                        set_cell_background(cell, OFF_WHITE_HEX)

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    build()
